from datetime import datetime
import io
import os
import re
import zipfile
import barcode
from barcode.writer import ImageWriter
import pandas as pd
import streamlit as st
from PIL import Image, ImageDraw, ImageFont
import openpyxl
from openpyxl.drawing.image import Image as OpenpyxlImage

st.set_page_config(
    page_title='STBバーコード自動生成・エクセル埋め込みツール', page_icon='📦', layout='centered'
)

st.title('📦 STBバーコード原本自動埋め込みツール')
st.write(
    'CSVをアップロードすると、ファイル名から機種名を自動抽出し、原本エクセルのレイアウトにバーコードを自動配置した印刷用エクセルファイルを生成します。'
)

# 1. CSVファイルのアップロード
uploaded_file = st.file_uploader(
    'STBリストのCSVファイルを選択してください', type=['csv']
)

if uploaded_file is not None:
  # ファイル名から機種名を自動抽出（例: "TZ-LS200P49台.csv" -> "TZ-LS200P"）
  filename_raw = uploaded_file.name
  extracted_model = "TZ-MODEL"
  
  match = re.match(r"^(.+?)(?:\d+台|\d+件|\.csv)", filename_raw)
  if match:
    extracted_model = match.group(1).strip()
  else:
    extracted_model = os.path.splitext(filename_raw)[0]
    extracted_model = re.sub(r'\d+.*$', '', extracted_model).strip()
    if not extracted_model:
      extracted_model = os.path.splitext(filename_raw)[0]

  df = pd.read_csv(uploaded_file, header=None)
  target_column = df.columns[0]

  # データの抽出と指数表記の自動修復
  cleaned_data_list = []
  for raw_data in df[target_column]:
    try:
      if isinstance(raw_data, float) or (
          isinstance(raw_data, str) and 'e' in raw_data.lower()
      ):
        val = str(int(float(raw_data)))
      else:
        val = str(raw_data).strip()
    except Exception:
      val = str(raw_data).strip()

    if val and val.lower() != 'nan':
      cleaned_data_list.append(val)

  st.success(
      f'✨ CSVから **{len(cleaned_data_list)}件** のデータを読み込みました！（自動抽出された機種名: **{extracted_model}**）'
  )

  # 2. 設定項目
  st.subheader('⚙️ バーコード設定')
  model_name_input = st.text_input(
      'バーコード・エクセル上に表示する機種名（自動抽出・編集可能）',
      value=extracted_model
  )

  # フォント取得ヘルパー
  def get_proper_font(size):
    current_dir = os.path.dirname(os.path.abspath(__file__))
    custom_font = os.path.join(current_dir, 'arial.ttf')
    font_paths = [
        custom_font,
        'C:/Windows/Fonts/meiryo.ttc',
        'C:/Windows/Fonts/YuGothM.ttc',
        'C:/Windows/Fonts/arial.ttf',
        '/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf',
        '/Library/Fonts/Arial.ttf'
    ]
    for path in font_paths:
      if os.path.exists(path):
        try:
          return ImageFont.truetype(path, size=int(size))
        except Exception:
          continue
    return ImageFont.load_default()

  # 1個分のラベル画像を生成する関数（機種名＋バーコード＋下部テキスト）
  def generate_single_label_image(clean_data, model_name):
    code39 = barcode.get_barcode_class('code39')
    barcode_instance = code39(clean_data, writer=ImageWriter(), add_checksum=False)

    options = {
        'module_width': 0.4,
        'module_height': 12.0,
        'quiet_zone': 8.0,
        'write_text': False,
    }

    rv = io.BytesIO()
    barcode_instance.write(rv, options=options)
    rv.seek(0)

    bc_img = Image.open(rv).convert('RGB')
    bc_w, bc_h = bc_img.size

    display_text = f"* {' '.join(list(clean_data))} *"
    font_size = max(18, int(bc_w * 0.045))
    font = get_proper_font(font_size)

    dummy_draw = ImageDraw.Draw(bc_img)
    try:
      char_widths = [dummy_draw.textlength(char, font=font) for char in display_text]
    except AttributeError:
      char_widths = [font.getlength(char) for char in display_text]

    sum_widths = sum(char_widths)
    left_margin = int(bc_w * 0.03)
    right_margin = int(bc_w * 0.03)
    available_width = bc_w - (left_margin + right_margin)

    if len(display_text) > 1:
      spacing = max(2, (available_width - sum_widths) / (len(display_text) - 1))
    else:
      spacing = 0

    header_h = 42
    footer_h = 35
    total_h = header_h + bc_h + footer_h
    
    label_img = Image.new('RGB', (bc_w, total_h), 'white')
    draw = ImageDraw.Draw(label_img)

    # 機種名（左上・大きくハッキリ）
    model_font_size = max(22, int(bc_w * 0.055))
    model_font = get_proper_font(model_font_size)
    for dx in [0, 1]:
      draw.text((left_margin + dx, 8), model_name, fill='black', font=model_font)

    # バーコード貼り付け
    label_img.paste(bc_img, (0, header_h))

    # 下部テキスト
    text_y = header_h + bc_h + 4
    current_x = left_margin
    for idx, char in enumerate(display_text):
      for dx in [0, 1]:
        draw.text((current_x + dx, text_y), char, fill='black', font=font)
      current_x += char_widths[idx] + spacing

    return label_img


  # 3. エクセル一括生成処理
  st.markdown('---')
  if cleaned_data_list:
    if st.button('📦 原本エクセルにバーコードを自動埋め込んで生成'):
      template_path = "STBﾊﾞｰｺｰﾄﾞ(620PW)原本.xlsx"
      
      if not os.path.exists(template_path):
        st.error(f"原本ファイル '{template_path}' が見つかりません。同じフォルダに配置してください。")
      else:
        wb = openpyxl.load_workbook(template_path)
        sheet_name = wb.sheetnames[0]
        ws = wb[sheet_name]

        # 原本のセル配置ルールに基づいてバーコードを埋め込む
        # 原本では5行おき（row 1, 6, 11...）に機種名があり、その下に画像が配置される構造
        # ここではcleaned_data_listの件数分、順番にセルへ機種名と画像をセットしていく
        
        # サンプルとして、全データを順次セルに配置していくロジック
        current_row = 1
        for idx, clean_data in enumerate(cleaned_data_list, start=1):
          # 1個分のバーコード画像を生成
          pil_img = generate_single_label_image(clean_data, model_name_input)
          temp_img_path = f"temp_gen_{idx}.png"
          pil_img.save(temp_img_path, format="PNG")

          # エクセル上の配置位置を計算（原本の構造に合わせて配置）
          # 例: 5行ごとにブロックが繰り返される場合
          block_row = ((idx - 1) // 2) * 5 + 1  # 2列構成などの場合に対応
          col_idx = 1 if (idx % 2 != 0) else 7  # A列かG列かなど

          # セルに機種名を設定
          ws.cell(row=block_row, column=col_idx).value = model_name_input

          # 画像を貼り付け
          img = OpenpyxlImage(temp_img_path)
          img.width = 220
          img.height = 80
          
          cell_coord = f"{openpyxl.utils.get_column_letter(col_idx)}{block_row + 1}"
          ws.add_image(img, cell_coord)

          if os.path.exists(temp_img_path):
            os.remove(temp_img_path)

        # 保存用バッファ
        output_buffer = io.BytesIO()
        wb.save(output_buffer)
        output_buffer.seek(0)

        date_str = datetime.now().strftime('%Y-%m-%d')
        dl_filename = f'STB_Barcodes_{model_name_input}_{date_str}.xlsx'

        st.success(f'✨ 全 {len(cleaned_data_list)}件のバーコードを原本エクセルに自動埋め込みしました！')
        st.download_button(
            label='📥 完成版エクセルファイルをダウンロード',
            data=output_buffer,
            file_name=dl_filename,
            mime='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        )
